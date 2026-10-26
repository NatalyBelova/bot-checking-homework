"""Таблицы для проверки чек-листов: скачивание по ссылке и перевод в текст для ИИ.

Поддерживаются:
- Google Таблицы (доступ «Все, у кого есть ссылка»), включая опубликованные
- Яндекс.Диск (публичная ссылка на файл)
- OneDrive личный (1drv.ms, onedrive.live.com) и SharePoint (если ссылка открыта без входа)
- прямые ссылки на .xlsx / .csv
- файлы .xlsx / .csv, присланные в Telegram
"""
import base64
import csv
import io
import re
import zipfile
from dataclasses import dataclass
from typing import List, Optional
from urllib.parse import parse_qsl, urlencode, urlparse, urlunparse

import aiohttp
from openpyxl import load_workbook

MAX_DOWNLOAD = 15 * 1024 * 1024   # 15 МБ
MAX_ROWS_PER_SHEET = 400
MAX_CELL = 200
MAX_TOTAL_CHARS = 40_000          # ≈ 12 тыс. токенов — по цене это ещё центы

URL_RE = re.compile(r"https?://[^\s<>\"'«»)\]]+")
SPREADSHEET_EXT = (".xlsx", ".xlsm", ".csv")


class TableError(Exception):
    """Понятная пользователю причина, почему таблицу не удалось получить."""


@dataclass
class Table:
    source: str  # ссылка или имя файла
    text: str    # содержимое для ИИ


# ---------------------------------------------------------------------------
# Ссылки
# ---------------------------------------------------------------------------

def find_urls(text):
    urls = []
    for url in URL_RE.findall(text or ""):
        url = url.rstrip(".,;:!?")
        if url not in urls:
            urls.append(url)
    return urls


def classify_url(url):
    """Куда ведёт ссылка: 'google' | 'yandex' | 'onedrive' | 'sharepoint' | 'direct' | None."""
    parsed = urlparse(url)
    host = parsed.netloc.lower()
    path = parsed.path.lower()
    if host.endswith("docs.google.com") and "/spreadsheets/" in path:
        return "google"
    if host in ("disk.yandex.ru", "disk.yandex.com", "yadi.sk") or host.endswith(".disk.yandex.ru"):
        return "yandex"
    if host in ("1drv.ms", "onedrive.live.com"):
        return "onedrive"
    if host.endswith(".sharepoint.com"):
        return "sharepoint"
    if path.endswith(SPREADSHEET_EXT):
        return "direct"
    return None


def google_export_url(url):
    parsed = urlparse(url)
    published = re.search(r"/spreadsheets/d/e/([\w-]+)", parsed.path)
    if published:
        return f"https://docs.google.com/spreadsheets/d/e/{published.group(1)}/pub?output=xlsx"
    match = re.search(r"/spreadsheets/d/([\w-]+)", parsed.path)
    if not match:
        raise TableError("не похоже на ссылку на Google Таблицу")
    return f"https://docs.google.com/spreadsheets/d/{match.group(1)}/export?format=xlsx"


def onedrive_download_url(url):
    token = base64.urlsafe_b64encode(url.encode()).decode().rstrip("=")
    return f"https://api.onedrive.com/v1.0/shares/u!{token}/root/content"


def sharepoint_download_url(url):
    parsed = urlparse(url)
    query = dict(parse_qsl(parsed.query))
    query["download"] = "1"
    return urlunparse(parsed._replace(query=urlencode(query)))


async def _download(session, url):
    async with session.get(url, allow_redirects=True) as response:
        if response.status in (401, 403):
            raise TableError("доступ закрыт")
        if response.status == 404:
            raise TableError("файл не найден")
        if response.status >= 400:
            raise TableError(f"сервер ответил ошибкой {response.status}")
        data = await response.content.read(MAX_DOWNLOAD + 1)
        if len(data) > MAX_DOWNLOAD:
            raise TableError("файл больше 15 МБ")
        return data, response.headers.get("Content-Type", ""), str(response.url)


async def _yandex_href(session, url):
    api = "https://cloud-api.yandex.net/v1/disk/public/resources/download"
    async with session.get(api, params={"public_key": url}) as response:
        if response.status == 404:
            raise TableError("файл на Яндекс.Диске не найден или доступ закрыт")
        if response.status >= 400:
            raise TableError(f"Яндекс.Диск ответил ошибкой {response.status}")
        return (await response.json())["href"]


async def fetch_url(url):
    """Скачивает таблицу по ссылке. Возвращает (bytes, filename). Бросает TableError."""
    kind = classify_url(url)
    if kind is None:
        raise TableError("это не ссылка на таблицу")

    timeout = aiohttp.ClientTimeout(total=40)
    headers = {"User-Agent": "Mozilla/5.0 (homework-bot)"}
    async with aiohttp.ClientSession(timeout=timeout, headers=headers) as session:
        try:
            if kind == "google":
                data, ctype, final = await _download(session, google_export_url(url))
            elif kind == "yandex":
                data, ctype, final = await _download(session, await _yandex_href(session, url))
            elif kind == "onedrive":
                data, ctype, final = await _download(session, onedrive_download_url(url))
            elif kind == "sharepoint":
                data, ctype, final = await _download(session, sharepoint_download_url(url))
            else:
                data, ctype, final = await _download(session, url)
        except aiohttp.ClientError as e:
            raise TableError(f"не удалось скачать ({type(e).__name__})")
        except TimeoutError:
            raise TableError("сервер слишком долго не отвечал")

    # закрытая Google-таблица отдаёт страницу входа (HTML) с кодом 200
    if data[:2] != b"PK" and ("text/html" in ctype or data.lstrip()[:1] == b"<"):
        if "accounts.google.com" in final or kind == "google":
            raise TableError("доступ закрыт — нужно открыть «Все, у кого есть ссылка»")
        raise TableError("по ссылке не таблица, а веб-страница (возможно, нужен вход)")

    name = "table.csv" if url.lower().endswith(".csv") or "text/csv" in ctype else "table.xlsx"
    return data, name


# ---------------------------------------------------------------------------
# Разбор файла в текст
# ---------------------------------------------------------------------------

def _cell(value):
    if value is None:
        return ""
    if hasattr(value, "strftime"):
        return value.strftime("%d.%m.%Y")
    text = str(value).replace("\n", " ").strip()
    return text if len(text) <= MAX_CELL else text[:MAX_CELL] + "…"


def _rows_to_text(title, rows):
    """Номера строк — как в самой таблице, чтобы «в строке 12» совпадало с тем, что видит ученик."""
    numbered = [(n, [_cell(v) for v in row]) for n, row in enumerate(rows, start=1)]
    numbered = [(n, r) for n, r in numbered if any(r)]
    if not numbered:
        return f"Лист «{title}»: пустой"
    width = max(max((i + 1 for i, v in enumerate(r) if v), default=0) for _, r in numbered)
    lines = [f"Лист «{title}» (непустых строк: {len(numbered)}; номер — строка в таблице):"]
    for n, row in numbered[:MAX_ROWS_PER_SHEET]:
        lines.append(f"{n}. " + " | ".join(row[:width]))
    if len(numbered) > MAX_ROWS_PER_SHEET:
        lines.append(f"… ещё {len(numbered) - MAX_ROWS_PER_SHEET} строк не показано")
    return "\n".join(lines)


def table_to_text(data, filename="table.xlsx"):
    """xlsx/csv → текст. Бросает TableError, если файл не читается."""
    name = filename.lower()
    if name.endswith(".csv"):
        for encoding in ("utf-8-sig", "cp1251"):
            try:
                decoded = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise TableError("не удалось прочитать CSV (неизвестная кодировка)")
        dialect = csv.Sniffer().sniff(decoded[:4000], delimiters=",;\t") if decoded.strip() else csv.excel
        text = _rows_to_text("CSV", csv.reader(io.StringIO(decoded), dialect))
    else:
        if name.endswith(".xls"):
            raise TableError("старый формат .xls не поддерживается — сохрани как .xlsx")
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except (zipfile.BadZipFile, KeyError, ValueError, OSError):
            raise TableError("файл не открывается как таблица Excel")
        parts = []
        for ws in wb.worksheets:
            if getattr(ws, "sheet_state", "visible") != "visible":
                continue
            parts.append(_rows_to_text(ws.title, ws.iter_rows(values_only=True)))
        wb.close()
        text = "\n\n".join(parts)

    if len(text) > MAX_TOTAL_CHARS:
        text = text[:MAX_TOTAL_CHARS] + "\n… таблица обрезана, показано начало"
    return text


def is_spreadsheet_name(path):
    return (path or "").lower().endswith(SPREADSHEET_EXT + (".xls",))


# ---------------------------------------------------------------------------
# Всё вместе
# ---------------------------------------------------------------------------

async def collect_tables(text, documents):
    """Таблицы из ссылок в тексте и из присланных файлов.

    documents — список (filename, bytes) уже скачанных из Telegram файлов-таблиц.
    Возвращает (tables, problems): problems — понятные причины по каждой неудаче.
    """
    tables: List[Table] = []
    problems: List[str] = []

    for url in find_urls(text):
        if classify_url(url) is None:
            continue
        try:
            data, name = await fetch_url(url)
            tables.append(Table(source=url, text=table_to_text(data, name)))
        except TableError as e:
            problems.append(f"{url} — {e}")

    for filename, data in documents:
        try:
            tables.append(Table(source=filename, text=table_to_text(data, filename)))
        except TableError as e:
            problems.append(f"{filename} — {e}")

    return tables, problems


def describe(tables: List[Table], problems: List[str], header: str) -> Optional[str]:
    if not tables and not problems:
        return None
    parts = []
    for i, t in enumerate(tables, start=1):
        parts.append(f"{header} №{i} (источник: {t.source}):\n<<<\n{t.text}\n>>>")
    if problems:
        parts.append("Не удалось открыть: " + "; ".join(problems))
    return "\n\n".join(parts)
