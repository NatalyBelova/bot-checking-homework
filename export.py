"""Выгрузка ДЗ в Excel."""
from datetime import timezone
from io import BytesIO
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from config import TIMEZONE

_TZ = ZoneInfo(TIMEZONE)

STATUS_TEXT = {"new": "На проверке", "revision": "На доработке", "accepted": "Принято"}
STATUS_FILL = {
    "accepted": PatternFill("solid", fgColor="D9EAD3"),
    "revision": PatternFill("solid", fgColor="FCE5CD"),
    "new": PatternFill("solid", fgColor="DDEBF7"),
}
HEADER_FILL = PatternFill("solid", fgColor="E7E6E6")
DATE_FORMAT = "DD.MM.YYYY HH:MM"


def _local(dt):
    """Excel не умеет часовые пояса: переводим UTC в местное время и убираем зону."""
    if dt is None:
        return None
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=timezone.utc)
    return dt.astimezone(_TZ).replace(tzinfo=None)


def _student(row):
    return row["student_name"] or f"id:{row['student_id']}"


def _reviewer(value):
    if not value:
        return ""
    return "ИИ" if value == "ai" else "Проверяющий"


def _sheet(wb, title, headers, widths, first=False):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True)
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"
    return ws


def build_workbook(assignments, homeworks, versions):
    wb = Workbook()

    # --- Прогресс: ученики × задания (по последнему ДЗ ученика на задание) ---
    titles = [a["title"] for a in assignments]
    ws = _sheet(wb, "Прогресс", ["Ученик", "ID"] + titles + ["Принято", "Всего заданий"],
                [24, 12] + [18] * len(titles) + [10, 10], first=True)

    students = {}
    for hw in homeworks:
        s = students.setdefault(hw["student_id"], {"name": _student(hw), "tasks": {}})
        if hw["assignment_id"] is not None:
            s["tasks"][hw["assignment_id"]] = hw  # ДЗ отсортированы по id → остаётся последнее

    for student_id, s in sorted(students.items(), key=lambda kv: kv[1]["name"].lower()):
        row = [s["name"], student_id]
        fills = []
        for a in assignments:
            hw = s["tasks"].get(a["id"])
            if hw:
                label = STATUS_TEXT.get(hw["status"], hw["status"])
                if hw["current_version"] > 1:
                    label += f" ({hw['current_version']})"
                row.append(label)
                fills.append(STATUS_FILL.get(hw["status"]))
            else:
                row.append("")
                fills.append(None)
        accepted = sum(1 for hw in s["tasks"].values() if hw["status"] == "accepted")
        row += [accepted, len(assignments)]
        ws.append(row)
        for i, fill in enumerate(fills):
            if fill:
                ws.cell(row=ws.max_row, column=3 + i).fill = fill
    if titles:
        ws.cell(row=ws.max_row + 2, column=1,
                value="В скобках — число попыток. Цвет: зелёный — принято, оранжевый — доработка, "
                      "голубой — на проверке.").font = Font(italic=True, color="808080")

    # --- Все ДЗ ---
    ws = _sheet(wb, "Все ДЗ",
                ["№ ДЗ", "Ученик", "ID ученика", "Задание", "Статус", "Попыток",
                 "Отправлено", "Обновлено", "Последний комментарий"],
                [8, 22, 12, 28, 14, 9, 17, 17, 60])
    for hw in homeworks:
        ws.append([
            hw["id"], _student(hw), hw["student_id"], hw["assignment_title"] or "",
            STATUS_TEXT.get(hw["status"], hw["status"]), hw["current_version"],
            _local(hw["created_at"]), _local(hw["updated_at"]), hw["last_comment"] or "",
        ])
        r = ws.max_row
        ws.cell(row=r, column=5).fill = STATUS_FILL.get(hw["status"], PatternFill())
        for col in (7, 8):
            ws.cell(row=r, column=col).number_format = DATE_FORMAT
        ws.cell(row=r, column=9).alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = ws.dimensions

    # --- Попытки ---
    ws = _sheet(wb, "Попытки",
                ["№ ДЗ", "Попытка", "Дата", "Текст ученика", "Файлов", "Комментарий", "Кто проверил"],
                [8, 9, 17, 50, 8, 60, 14])
    for v in versions:
        ws.append([
            v["homework_id"], v["version_num"], _local(v["created_at"]), v["text"] or "",
            v["files"], v["comment"] or "", _reviewer(v["reviewed_by"]) if v["comment"] is not None else "",
        ])
        r = ws.max_row
        ws.cell(row=r, column=3).number_format = DATE_FORMAT
        for col in (4, 6):
            ws.cell(row=r, column=col).alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = ws.dimensions

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
